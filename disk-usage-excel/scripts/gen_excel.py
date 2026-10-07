# -*- coding: utf-8 -*-
"""disk-usage-excel: 将 collect_sizes.ps1 生成的 JSON 渲染为 Excel 报表。

输出路径可显式传入（argv[2]），留空则从输入 JSON 所在目录动态推导，
避免写死个人工作区路径。
"""
import json
import sys
import io
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_FILL = PatternFill("solid", fgColor="FFF2CC")


def make_sheet(ws, title, headers, rows, highlight_first_col_prefix=None):
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
        if highlight_first_col_prefix and str(row[0].value or "").startswith(highlight_first_col_prefix):
            for c in row:
                c.fill = TOP_FILL
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        maxlen = 0
        for cell in col:
            v = str(cell.value) if cell.value is not None else ""
            w = sum(2 if ord(ch) > 127 else 1 for ch in v)
            maxlen = max(maxlen, w)
        ws.column_dimensions[get_column_letter(i)].width = min(maxlen + 4, 75)


def rows_of(items, root=None):
    out = []
    for it in items or []:
        p = it.get("Path", "")
        if root and p.startswith(root):
            p = p[len(root):].lstrip("\\")
        out.append([p, it.get("SizeGB", 0), it.get("FileCount", 0), it.get("SizeBytes")])
    return out


def root_tag(root):
    """把根目录压成短标签：盘符根用 C，非盘符根取最后一级目录名。"""
    label = (root or "").strip().rstrip("\\/")
    if not label:
        return "root"
    if re.fullmatch(r"[A-Za-z]:", label):
        return label[0].upper()
    if re.fullmatch(r"[A-Za-z]:\\", label):
        return label[0].upper()
    parts = [p for p in re.split(r"[\\/]", label) if p]
    return parts[-1] if parts else "root"


def default_out_path(json_path, root):
    """输出目录取 JSON 所在目录，文件名由 JSON 名 + 根目录标签推导，无写死路径。"""
    base = os.path.splitext(os.path.basename(json_path))[0]
    tag = root_tag(root)
    return os.path.join(
        os.path.dirname(os.path.abspath(json_path)),
        "{0}_{1}盘_磁盘空间统计.xlsx".format(base, tag),
    )


def main():
    if len(sys.argv) < 2:
        print("usage: gen_excel.py <disk_sizes.json> [output.xlsx]", file=sys.stderr)
        return 2
    json_path = sys.argv[1]
    with io.open(json_path, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    root_label = data.get("Root", "")
    out_xlsx = sys.argv[2] if len(sys.argv) > 2 and sys.argv[2].strip() else \
        default_out_path(json_path, root_label)

    root = root_label.rstrip("\\") + "\\"
    wb = Workbook()
    headers = ["路径", "大小(GB)", "文件数", "实际大小(Bytes)"]
    make_sheet(wb.active, "第一层", headers, rows_of(data.get("Level1")))

    ws2 = wb.create_sheet()
    make_sheet(ws2, "第二层", headers, rows_of(data.get("Level2")))

    deep = data.get("Deep") or []
    for depth, label in [(3, "第三层"), (4, "第四层"), (5, "第五层"), (6, "第六层")]:
        rows = []
        for detail in deep:
            rows.extend(row + [detail["TopPath"]]
                        for row in rows_of(detail.get("Level{0}".format(depth))))
        make_sheet(wb.create_sheet(), label + "(深度分析)", headers + ["所属Top目录"], rows)

    total = (round(data["Level1TotalBytes"] / (1024 ** 3), 2)
             if data.get("Level1TotalBytes") is not None else
             data.get("Validation", {}).get("Level1TotalGB",
                 round(sum(it.get("SizeGB", 0) for it in data.get("Level1") or []), 2)))
    val = data.get("Validation") or {}
    verdict = val.get("Verdict", "unknown")
    warn_rows = []
    if verdict == "suspicious-low":
        warn_rows = [
            ["数据完整性", "不完整 - 统计结果不可直接使用"],
            ["原因", val.get("Note", "")],
            ["建议", "取得完全访问权限后重新运行 collect_sizes.ps1"],
        ]
    elif verdict == "suspicious-high":
        warn_rows = [["数据完整性", "存疑 - " + val.get("Note", "")]]

    ws5 = wb.create_sheet()
    summary_rows = [
        ["统计根目录", root_label],
        ["生成时间", data.get("GeneratedAt", "")],
        ["第一层全量总大小(GB)", total],
    ]
    if val.get("Checked"):
        summary_rows += [
            ["该盘真实已用(GB)", val.get("DiskUsedGB")],
            ["覆盖率(%)", val.get("CoveragePct")],
        ]
    summary_rows += [
        ["第一层全量文件夹数", data.get("Level1FolderCount", len(data.get("Level1") or []))],
        ["第一层展示文件夹数", len(data.get("Level1") or [])],
        ["第二层展示文件夹数", len(data.get("Level2") or [])],
        ["筛选口径", "每个父目录分别保留 Top{0}，沿保留目录展开".format(data["PerParentLimit"])
         if data.get("PerParentLimit") else "旧版全量展示"],
        ["大小口径", "GB 四舍五入显示；父子大小合计请核对实际大小(Bytes)列"],
        ["深度分析目录", ", ".join(d["TopPath"] for d in deep)],
    ]
    make_sheet(ws5, "汇总", ["项目", "值"], summary_rows)

    if warn_rows:
        ws6 = wb.create_sheet()
        make_sheet(ws6, "数据完整性告警", ["项目", "说明"], warn_rows)
        ws6.column_dimensions["B"].width = 70
        for row in ws6.iter_rows(min_row=2):
            row[0].fill = PatternFill("solid", fgColor="FFC7CE")
            row[0].font = Font(color="9C0006", bold=True)

    wb.save(out_xlsx)
    print("saved:", out_xlsx)
    return 0


if __name__ == "__main__":
    sys.exit(main())
